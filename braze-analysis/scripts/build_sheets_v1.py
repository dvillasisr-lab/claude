"""Construye braze_auditoria_v1.xlsx: Léeme, Parámetros, 15 pestañas crudas y Verificación.
Uso: python3 scripts/build_sheets_v1.py   (desde braze-analysis/)
Las fórmulas usan sintaxis de Google Sheets; el archivo se abre con Google Sheets, no con Excel.
"""
import pandas as pd, hashlib, re, os, datetime
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment

D = 'data/'
OUT = 'sheets/braze_auditoria_v1.xlsx'
HOY = '14 de octubre de 2026'

# pestaña, archivo, qué contiene, grano, columna sin vacíos para contar filas
TABS = [
 ('E1',  'e1_comunicacion_mes_canal.csv',   'Volumen por comunicación × mes × canal: envíos, entregas, aperturas, clics, bajas, conversiones. Es la base del análisis por campaña.', 'comunicación × mes × canal (deduplicado por usuario × comunicación × día)', 'communication_id'),
 ('E2',  'e2_catalogo.csv',                 'Catálogo de piezas de Braze (campañas y canvas): nombre, equipo, producto, ticket, canales, fechas de primer y último envío, estado.', 'una fila por communication_id', 'communication_id'),
 ('E3',  'e3_usuarios_distribucion.csv',    'Distribución de usuarios por banda de mensajes recibidos en el periodo, suscripción, etapa y contactabilidad.', 'banda × suscripción × etapa (usuarios agregados)', 'banda_msgs'),
 ('E4a', 'e4a_serie_semana_mes_canal.csv',  'Serie de tiempo por canal con dos granos: semana y mes. Toda fórmula debe filtrar la columna grano.', 'grano × periodo × canal', 'grano'),
 ('E4b', 'e4b_concentracion_mes.csv',       'Concentración: cuántos usuarios recibieron cuántos mensajes cada mes (bandas).', 'mes × banda de mensajes', 'mes'),
 ('E5',  'e5_cohorte_baja.csv',             'Cohorte de usuarios dados de baja por mes y qué recibieron después de la baja (meses 1, 2, 3 y total).', 'mes de baja', 'mes_baja'),
 ('E5b', 'e5b_comms_despues_baja.csv',      'Familias que siguieron mandando email a usuarios ya dados de baja.', 'familia', 'familia'),
 ('E6a', 'e6a_solape_dia_mes.csv',          'Solapamiento: cuántas familias distintas recibe un usuario el mismo día, por mes.', 'mes × banda de familias por día', 'mes'),
 ('E6b', 'e6b_pares_familias.csv',          'Pares de familias que llegan al mismo usuario el mismo día (top 1,000 por LIMIT en la query).', 'par de familias', 'familia_a'),
 ('E7a', 'e7a_familia_canal_etapa.csv',     'Envíos por familia × canal × etapa del usuario (golden record: Blacklist, CC Freeze, Reus/Arco, Prospects, sin_golden).', 'familia × canal × etapa', 'canal'),
 ('E7b', 'e7b_patron_respuesta.csv',        'Patrón de respuesta por usuario en todo el periodo: 16 combinaciones de respondió o no por canal, con mensajes recibidos, conversiones y bajas.', '16 combinaciones de respuesta (usuarios agregados)', 'usuarios'),
 ('E7c', 'e7c_abren_vs_no_abren.csv',       'Por familia y canal: usuarios que abrieron o no, y cuántos convirtieron en cada grupo (todo el periodo).', 'familia × canal', 'canal'),
 ('E8',  'e8_dias_hasta_baja.csv',          'Días y emails desde el primer email hasta la baja. Bloques: resumen, banda_dias, banda_emails y familia_dia_baja (familia enviada el día de la baja). cohorte_limpia 1 = primer email dentro del periodo; 0 = ya recibía antes de enero.', 'bloque × clave × cohorte', 'bloque'),
 ('E9a', 'e9a_abren_vs_no_abren_ventanas.csv', 'Como E7c pero con ventanas de conversión a 7, 30, 90 días y fin de periodo.', 'familia × canal', 'canal'),
 ('E9b', 'e9b_patron_respuesta_90d.csv',    'Como E7b pero solo julio a septiembre 2026 (último trimestre).', '16 combinaciones de respuesta (usuarios agregados)', 'usuarios'),
]
TEXT_COLS = {'communication_id','ticket_number','banda_msgs','banda_msgs_mes','banda_familias_dia','clave','fecha_codigo','grupo','mes','periodo','mes_baja','first_sent','last_sent','created_at','tipo_pieza','nombre','familia','familia_a','familia_b'}
NUM_RE = re.compile(r'^-?\d+(\.\d+)?$')

def md5(p):
    return hashlib.md5(open(p,'rb').read()).hexdigest()

def read(path):
    return pd.read_csv(path, dtype=str, keep_default_na=False)

def typed(series, name):
    """Devuelve lista de valores con tipo y el conteo de 'null' convertidos a vacío."""
    vals = series.tolist(); nulls = 0; out = []
    nonnull = [v for v in vals if v not in ('', 'null')]
    numeric = name not in TEXT_COLS and len(nonnull) > 0 and all(NUM_RE.match(v) for v in nonnull)
    allint = numeric and all('.' not in v for v in nonnull)
    for v in vals:
        if v == 'null': nulls += 1; out.append(None)
        elif v == '': out.append(None)
        elif numeric: out.append(int(v) if allint else float(v))
        else: out.append(v)
    return out, nulls, ('número' if numeric else 'texto')

wb = Workbook()
wb.remove(wb.active)
bold = Font(bold=True)
INFO = {}   # por pestaña: filas, md5, columnas, letras, nulos
E = {}      # dataframes

def write_raw(tab, fname, derived):
    df = read(D+fname); E[tab] = df
    ws = wb.create_sheet(tab)
    cols = list(df.columns); letters = {}
    nulls = {}; tipos = {}
    data_cols = []
    for j, c in enumerate(cols, start=1):
        ws.cell(1, j, c).font = bold
        letters[c] = L(j)
        vals, n, t = typed(df[c], c); data_cols.append(vals); nulls[c] = n; tipos[c] = t
    n = len(df)
    for i in range(n):
        r = i + 2
        for j, vals in enumerate(data_cols, start=1):
            v = vals[i]
            if v is not None: ws.cell(r, j, v)
    # columnas derivadas (fórmula por fila)
    for k, (hdr, fn) in enumerate(derived):
        j = len(cols) + 1 + k
        ws.cell(1, j, hdr).font = bold
        letters[hdr] = L(j)
        for i in range(n):
            ws.cell(i+2, j, fn(i+2, letters))
    last = L(len(cols) + len(derived))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{last}{n+1}'
    for j in range(1, len(cols)+len(derived)+1):
        ws.column_dimensions[L(j)].width = 14
    INFO[tab] = dict(filas=n, md5=md5(D+fname), cols=cols, letters=letters, nulls=nulls, tipos=tipos, derived=[h for h,_ in derived], archivo=fname)
    return ws

DER = {
 'E1': [('familia (derivada)', lambda r,l: f'=IF({l["communication_name"]}{r}="","(sin nombre)-"&{l["communication_id"]}{r},REGEXREPLACE({l["communication_name"]}{r},"-\\d{{8}}-","-"))'),
        ('mes_clave (derivada)', lambda r,l: f'=LEFT({l["mes"]}{r},7)')],
 'E2': [('familia (derivada)', lambda r,l: f'=REGEXREPLACE({l["nombre"]}{r},"-\\d{{8}}-","-")')],
 'E4a':[('mes_clave (derivada)', lambda r,l: f'=LEFT({l["periodo"]}{r},7)')],
 'E4b':[('mes_clave (derivada)', lambda r,l: f'=LEFT({l["mes"]}{r},7)')],
 'E6a':[('mes_clave (derivada)', lambda r,l: f'=LEFT({l["mes"]}{r},7)')],
}

# ---------- Léeme y Parámetros se crean primero para que queden al inicio
leeme = wb.create_sheet('Léeme')
par = wb.create_sheet('Parámetros')
for tab, fname, *_ in TABS:
    write_raw(tab, fname, DER.get(tab, []))
ver = wb.create_sheet('Verificación')

# ---------- Parámetros
P = [
 ('parámetro','valor','unidad','nota'),
 ('Costo por email enviado', 0.05, 'MXN por email_sent', 'Dato del usuario. Aplica a todo email enviado.'),
 ('Costo por WhatsApp', 1.00, 'MXN por mensaje', 'Dato del usuario. Se multiplica por entregados o por enviados según la fila 6.'),
 ('Costo por push', 0, 'MXN por pn_send', 'Sin costo variable.'),
 ('Costo por in-app', 0, 'MXN por in_app_impression', 'Sin costo variable.'),
 ('Base de cobro WhatsApp', 'entregado', 'entregado | enviado', 'Meta cobra por mensaje entregado (default). Cambia a "enviado" para ver el escenario que se usó en las muestras anteriores. Toda fórmula de costo lee esta celda.'),
 ('Costo fijo Braze por mes', None, 'MXN', 'PENDIENTE del usuario: significado de $7,020/mes y 32,474 créditos. Vacío hasta tenerlo; ninguna fórmula lo suma mientras esté vacío.'),
 ('Usuarios con email en el periodo', 20137560, 'usuarios', 'Denominador de la tasa de baja (E5, línea de denominador). Count distinct aproximado.'),
 ('Periodo inicio', '2026-01-01', 'fecha', 'Rango de event_date en todas las queries.'),
 ('Periodo fin', '2026-09-30', 'fecha', ''),
 ('Mínimo de respuesta email (clic / entregado)', 0.003, 'proporción', 'Debajo de esto la campaña "no tiene señal" en email. Editable.'),
 ('Mínimo de respuesta push (apertura / enviado)', 0.0075, 'proporción', 'Editable.'),
 ('Mínimo de respuesta WhatsApp (lectura / entregado)', 0.60, 'proporción', 'Editable.'),
 ('Mínimo de respuesta in-app (clic / impresión)', 0.10, 'proporción', 'Editable.'),
 ('Fracción del mínimo para marcar "sin señal"', 0.5, 'proporción', 'Regla de acción: respuesta menor a (mínimo × esta fracción) = sin señal.'),
 ('Lift mínimo (convierten los que abren / los que no abren)', 1.5, 'veces', 'Regla "mueve": conv de los que abren ≥ mitad de la mediana del canal y lift ≥ este valor.'),
 ('Muestra mínima: usuarios que abrieron', 1000, 'usuarios', 'Debajo no se juzga la campaña por conversión.'),
 ('Muestra mínima: usuarios que no abrieron', 1000, 'usuarios', ''),
 ('Muestra mínima: conversiones', 100, 'conversiones', ''),
 ('In-app: clic debajo del cual se quita la pieza', 0.05, 'proporción', ''),
 ('Definición de respuesta: email', 'clic (email_click)', 'texto', 'No se usa la apertura de email porque Apple y Gmail abren automáticamente.'),
 ('Definición de respuesta: push', 'apertura (pn_open)', 'texto', ''),
 ('Definición de respuesta: WhatsApp', 'lectura (whatsapp_read)', 'texto', ''),
 ('Definición de respuesta: in-app', 'clic (in_app_click)', 'texto', ''),
 ('Definición de baja', 'email_unsubscribe = true', 'texto', 'Solo baja de email. Usuario-mes (E4a) = un usuario cuenta cada mes que se da de baja; usuarios únicos (E5, E7b) = una vez en el periodo.'),
]
for i, row in enumerate(P, start=1):
    for j, v in enumerate(row, start=1):
        c = par.cell(i, j, v)
        if i == 1: c.font = bold
for r in (11,12,13,14,15,20): par.cell(r,2).number_format = '0.00%'
par.cell(2,2).number_format = '0.00'; par.cell(3,2).number_format = '0.00'
dv = DataValidation(type='list', formula1='"entregado,enviado"', allow_blank=False); par.add_data_validation(dv); dv.add('B6')
par.column_dimensions['A'].width = 52; par.column_dimensions['B'].width = 16; par.column_dimensions['C'].width = 26; par.column_dimensions['D'].width = 110
par.freeze_panes = 'A2'
PB = {k: f'Parámetros!$B${r}' for k, r in dict(costo_email=2, costo_wa=3, costo_push=4, costo_inapp=5, base_wa=6, usuarios_email=8).items()}

# ---------- helpers para Verificación
def rng(tab, col):
    i = INFO[tab]; return f"{tab}!${i['letters'][col]}$2:${i['letters'][col]}${i['filas']+1}"
def num(tab, col):
    return pd.to_numeric(E[tab][col].replace({'null': None, '': None}), errors='coerce')
def S(tab, col): return float(num(tab, col).sum())
def fam_pd(s): return s.str.replace(r'-\d{8}-', '-', regex=True)

e1 = E['E1']; e1f = fam_pd(e1.communication_name.where(e1.communication_name != 'null', None))
e1f = e1f.fillna('(sin nombre)-' + e1.communication_id)
e1mes = e1.mes.str[:7]; e1env = num('E1','envios')
e4a = E['E4a']; mes4 = e4a.grano == 'mes'
def S4(col, g='mes'): return float(num('E4a', col)[e4a.grano == g].sum())
e8 = E['E8']
costo_ent = S('E1','email_sent')*0.05 + S('E1','whatsapp_delivery')*1.0
costo_env = S('E1','email_sent')*0.05 + S('E1','whatsapp_send')*1.0

V = []  # (grupo, pestaña, comprobación, fórmula, esperado, explicación)
def add(g, t, q, f, exp, ex=''): V.append((g, t, q, f, exp, ex))

for tab, *_ , keycol in TABS:
    add('Filas', tab, f'Filas de datos (cuenta {keycol}, columna sin vacíos)', f'=COUNTA({rng(tab, keycol)})', INFO[tab]['filas'], 'Debe ser igual al CSV. Si difiere, la pega quedó incompleta.')

# E1 totales
for c in ['envios','email_sent','email_delivery','email_open','email_click','email_unsubscribe','pn_send','pn_open','whatsapp_send','whatsapp_delivery','whatsapp_read','in_app_impression','in_app_click','banner_impression','banner_click','conversiones_filas','conversiones_fraccional','usuarios_aprox']:
    add('E1 totales', 'E1', f'Suma de {c}', f'=SUM({rng("E1", c)})', S('E1', c), 'usuarios_aprox suma usuarios por fila: no es usuarios únicos.' if c=='usuarios_aprox' else '')
add('E1 totales', 'E1', 'Suma de costo_mxn precalculado (NO USAR: WhatsApp sobre enviados)', f'=SUM({rng("E1","costo_mxn")})', S('E1','costo_mxn'), 'Columna heredada de la query. Coincide con el costo "base enviado" de abajo. El costo del análisis sale siempre de Parámetros.')
add('E1 totales', 'E1', 'Celdas vacías en monto_filas (eran null en Databricks)', f'=COUNTBLANK({rng("E1","monto_filas")})', int((e1.monto_filas=='null').sum()), 'null = sin conversiones con monto. Se dejó vacío para que SUM funcione.')
add('E1 totales', 'E1', 'Filas con envíos = 0', f'=COUNTIF({rng("E1","envios")},0)', int((e1env==0).sum()), 'Meses con aperturas o bajas de envíos anteriores pero sin envío nuevo. Las campañas activas se cuentan con envíos > 0.')
add('E1 totales', 'E1', 'Filas con ticket_number = "Other"', f'=COUNTIF({rng("E1","ticket_number")},"Other")', int((e1.ticket_number=='Other').sum()), 'Sin ticket en Braze. El Catálogo (v2) intenta sacarlo del nombre con regex.')
add('E1 totales', 'E1', 'Filas sin communication_name (una comunicación sin nombre, 3 meses)', f'=COUNTBLANK({rng("E1","communication_name")})', int((e1.communication_name=='null').sum()), 'Un solo communication_id, 325 mil emails en julio y agosto. En familia (derivada) aparece como "(sin nombre)-id".')
add('E1 derivadas', 'E1', 'communication_id distintos', f'=COUNTA(UNIQUE({rng("E1","communication_id")}))', int(e1.communication_id.nunique()), '')
add('E1 derivadas', 'E1', 'familias distintas (incluye "(sin nombre)")', f'=COUNTA(UNIQUE({rng("E1","familia (derivada)")}))', int(e1f.nunique()), 'familia = nombre sin la fecha -AAAAMMDD-. Comprueba que la fórmula REGEXREPLACE evaluó en todas las filas.')
add('E1 derivadas', 'E1', 'Filas con mes_clave con forma 2026-MM', f'=COUNTIF({rng("E1","mes_clave (derivada)")},"2026-*")', int((e1mes.str.startswith('2026')).sum()), 'Debe ser igual a las filas. Si no, alguna fecha llegó en otro formato.')
for m in sorted(e1mes.unique()):
    add('E1 derivadas', 'E1', f'Filas del mes {m}', f'=COUNTIF({rng("E1","mes_clave (derivada)")},"{m}")', int((e1mes==m).sum()), '')
add('E1 derivadas', 'E1', 'Filas de septiembre con envíos > 0', f'=COUNTIFS({rng("E1","mes_clave (derivada)")},"2026-09",{rng("E1","envios")},">0")', int(((e1mes=='2026-09')&(e1env>0)).sum()), '')
add('E1 derivadas', 'E1', 'Familias activas en septiembre (envíos > 0)', f'=COUNTA(UNIQUE(FILTER({rng("E1","familia (derivada)")},{rng("E1","mes_clave (derivada)")}="2026-09",{rng("E1","envios")}>0)))', int(e1f[(e1mes=='2026-09')&(e1env>0)].nunique()), '')
# Costos
add('Costos', 'E1 + Parámetros', 'Costo variable total con la base de cobro seleccionada', f'=SUM({rng("E1","email_sent")})*{PB["costo_email"]}+IF({PB["base_wa"]}="entregado",SUM({rng("E1","whatsapp_delivery")}),SUM({rng("E1","whatsapp_send")}))*{PB["costo_wa"]}+SUM({rng("E1","pn_send")})*{PB["costo_push"]}+SUM({rng("E1","in_app_impression")})*{PB["costo_inapp"]}', f'=IF({PB["base_wa"]}="entregado",{costo_ent},{costo_env})', 'Prueba del selector: cambia Parámetros!B6 entre entregado y enviado; esta fila y la esperada deben moverse juntas.')
add('Costos', 'E1', 'Costo variable si WhatsApp se cobra por entregado (fijo, para comparar)', f'=SUM({rng("E1","email_sent")})*0.05+SUM({rng("E1","whatsapp_delivery")})*1', costo_ent, 'Supone 0.05 por email y 1.00 por WhatsApp. Es la cifra de las muestras actualizadas (MUESTRA_J).')
add('Costos', 'E1', 'Costo variable si WhatsApp se cobra por enviado (fijo, para comparar)', f'=SUM({rng("E1","email_sent")})*0.05+SUM({rng("E1","whatsapp_send")})*1', costo_env, 'Cifra de las muestras viejas (MUESTRA_B y C: 71.9 M). Queda superada.')
add('Costos', 'E1', 'WhatsApp enviados que no se entregaron', f'=SUM({rng("E1","whatsapp_send")})-SUM({rng("E1","whatsapp_delivery")})', S('E1','whatsapp_send')-S('E1','whatsapp_delivery'), '14.2 M mensajes. Con base entregado no cuestan; con base enviado costarían 14.2 M MXN.')
# E4a
add('E4a vs E1', 'E4a', 'Filas con grano = mes', f'=COUNTIF({rng("E4a","grano")},"mes")', int(mes4.sum()), '9 meses × 5 canales.')
add('E4a vs E1', 'E4a', 'Filas con grano = semana', f'=COUNTIF({rng("E4a","grano")},"semana")', int((~mes4).sum()), '')
for c, ex in [('envios','E4a cuenta eventos; E1 deduplica usuario × comunicación × día. Diferencia esperada menor a 0.01 %.'),('email_sent',''),('pn_send',''),('whatsapp_send',''),('whatsapp_delivery',''),('in_app_impression','Misma razón que envíos.'),('email_click',''),('bajas_email_eventos','Igual a email_unsubscribe de E1 (eventos).')]:
    e1c = 'email_unsubscribe' if c=='bajas_email_eventos' else c
    add('E4a vs E1', 'E4a', f'Suma de {c} con grano = mes', f'=SUMIF({rng("E4a","grano")},"mes",{rng("E4a",c)})', S4(c), ex)
    add('E4a vs E1', 'E4a', f'  diferencia contra E1 ({e1c})', f'=SUMIF({rng("E4a","grano")},"mes",{rng("E4a",c)})-SUM({rng("E1",e1c)})', S4(c)-S('E1',e1c), 'Cero o casi cero.' if not ex else ex)
add('Bajas', 'E4a', 'Bajas usuario-mes (bajas_email_usuarios, grano mes)', f'=SUMIF({rng("E4a","grano")},"mes",{rng("E4a","bajas_email_usuarios")})', S4('bajas_email_usuarios'), 'CIFRA A: 457,376. Un usuario cuenta en cada mes en que se dio de baja.')
add('Bajas', 'E4a', 'Bajas usuario-semana (grano semana)', f'=SUMIF({rng("E4a","grano")},"semana",{rng("E4a","bajas_email_usuarios")})', S4('bajas_email_usuarios','semana'), 'Mayor que usuario-mes porque un usuario puede caer en dos semanas del mismo mes.')
add('Bajas', 'E5', 'Usuarios únicos con baja (suma usuarios_baja)', f'=SUM({rng("E5","usuarios_baja")})', S('E5','usuarios_baja'), 'CIFRA B: 439,269. Un usuario cuenta una vez en el periodo.')
add('Bajas', 'E7b', 'Usuarios únicos con baja (suma usuarios_con_baja_email)', f'=SUM({rng("E7b","usuarios_con_baja_email")})', S('E7b','usuarios_con_baja_email'), '')
add('Bajas', 'E5 vs E7b', 'E5 menos E7b (debe dar cero)', f'=SUM({rng("E5","usuarios_baja")})-SUM({rng("E7b","usuarios_con_baja_email")})', S('E5','usuarios_baja')-S('E7b','usuarios_con_baja_email'), 'Dos queries distintas, misma definición.')
add('Bajas', 'E4a vs E5', 'Usuario-mes menos únicos', f'=SUMIF({rng("E4a","grano")},"mes",{rng("E4a","bajas_email_usuarios")})-SUM({rng("E5","usuarios_baja")})', S4('bajas_email_usuarios')-S('E5','usuarios_baja'), '18 mil usuarios se dieron de baja en más de un mes. Las dos cifras son correctas con su definición.')
add('Bajas', 'E1', 'Eventos de baja (email_unsubscribe)', f'=SUM({rng("E1","email_unsubscribe")})', S('E1','email_unsubscribe'), '1.09 M eventos: un usuario puede generar varios eventos de baja en un mes. No usar como usuarios.')
add('Bajas', 'E3', 'Eventos de baja según E3 (bajas_email)', f'=SUM({rng("E3","bajas_email")})', S('E3','bajas_email'), '0.2 % arriba de E1 por deduplicación de E1.')
add('Bajas', 'E4b', 'usuarios_con_baja_email (usuario-mes dentro de banda)', f'=SUM({rng("E4b","usuarios_con_baja_email")})', S('E4b','usuarios_con_baja_email'), '1.3 % debajo de E4a. Diferencia no explicada del todo (conteo aproximado y usuarios sin fila en la banda). E4b no se usa para contar bajas, solo para concentración de mensajes.')
add('Bajas', 'E5 + Parámetros', 'Tasa de baja del periodo = únicos / usuarios con email', f'=SUM({rng("E5","usuarios_baja")})/{PB["usuarios_email"]}', S('E5','usuarios_baja')/20137560, '2.2 %. Denominador en Parámetros.')
# E3 / E7b
add('E3 vs E7b', 'E3', 'Usuarios (suma)', f'=SUM({rng("E3","usuarios")})', S('E3','usuarios'), '19.5 M usuarios con al menos una fila en el periodo.')
add('E3 vs E7b', 'E7b', 'Usuarios (suma)', f'=SUM({rng("E7b","usuarios")})', S('E7b','usuarios'), 'Debe ser igual a E3.')
for c3, c7 in [('msgs','msgs'),('msgs_email','msgs_email'),('msgs_push','msgs_push'),('msgs_whatsapp','msgs_wa'),('msgs_whatsapp_entregados','msgs_wa_entregados'),('msgs_inapp','msgs_inapp')]:
    f3 = f'=SUM({rng("E3",c3)})'
    if c7 == 'msgs': f7 = f'=SUM({rng("E7b","msgs_email")})+SUM({rng("E7b","msgs_push")})+SUM({rng("E7b","msgs_wa")})+SUM({rng("E7b","msgs_inapp")})'; v7 = S('E7b','msgs_email')+S('E7b','msgs_push')+S('E7b','msgs_wa')+S('E7b','msgs_inapp')
    else: f7 = f'=SUM({rng("E7b",c7)})'; v7 = S('E7b',c7)
    add('E3 vs E7b', 'E3', f'Suma de {c3}', f3, S('E3',c3), '')
    add('E3 vs E7b', 'E7b', f'Suma de {c7} (E7b menos E3, debe dar cero)', f'{f7}-SUM({rng("E3",c3)})', v7-S('E3',c3), 'Mismo grano (filas por usuario).')
add('E3 vs E1', 'E3', 'msgs_email de E3 menos email_sent de E1', f'=SUM({rng("E3","msgs_email")})-SUM({rng("E1","email_sent")})', S('E3','msgs_email')-S('E1','email_sent'), 'E3 cuenta filas; E1 deduplica usuario × comunicación × día. 0.2 % esperado.')
add('E3 vs E1', 'E3', 'msgs_whatsapp_entregados de E3 menos whatsapp_delivery de E1', f'=SUM({rng("E3","msgs_whatsapp_entregados")})-SUM({rng("E1","whatsapp_delivery")})', S('E3','msgs_whatsapp_entregados')-S('E1','whatsapp_delivery'), '1.3 % esperado por la misma razón.')
add('E3 vs E1', 'E4b', 'msgs de E4b menos msgs de E3 (debe dar cero)', f'=SUM({rng("E4b","msgs")})-SUM({rng("E3","msgs")})', S('E4b','msgs')-S('E3','msgs'), '')
add('E7b', 'E7b', 'Usuarios que nunca respondieron (fila con los cuatro resp en 0)', f'=SUMIFS({rng("E7b","usuarios")},{rng("E7b","resp_email")},0,{rng("E7b","resp_push")},0,{rng("E7b","resp_wa")},0,{rng("E7b","resp_inapp")},0)', float(num('E7b','usuarios')[(E['E7b'].resp_email=='0')&(E['E7b'].resp_push=='0')&(E['E7b'].resp_wa=='0')&(E['E7b'].resp_inapp=='0')].sum()), '14.2 M de 19.5 M (73 %). Ver definición explícita en Léeme.')
add('E7b', 'E7b', 'Conversiones de usuarios (usuarios_con_conversion)', f'=SUM({rng("E7b","usuarios_con_conversion")})', S('E7b','usuarios_con_conversion'), '1.19 M usuarios. No comparar con conversiones_fraccional de E1 (3.38 M, por fila y fraccional).')
# E9b
add('E9b', 'E9b', 'Usuarios jul-sep (suma)', f'=SUM({rng("E9b","usuarios")})', S('E9b','usuarios'), '7.07 M usuarios con mensajes en el trimestre.')
add('E9b', 'E9b', 'Usuarios que nunca respondieron en jul-sep', f'=SUMIFS({rng("E9b","usuarios")},{rng("E9b","resp_email")},0,{rng("E9b","resp_push")},0,{rng("E9b","resp_wa")},0,{rng("E9b","resp_inapp")},0)', float(num('E9b','usuarios')[(E['E9b'].resp_email=='0')&(E['E9b'].resp_push=='0')&(E['E9b'].resp_wa=='0')&(E['E9b'].resp_inapp=='0')].sum()), '4.10 M (58 %).')
add('E9b', 'E9b vs E1', 'msgs_email de E9b menos email_sent de E1 jul-sep', f'=SUM({rng("E9b","msgs_email")})-SUMIF({rng("E1","mes_clave (derivada)")},"2026-07",{rng("E1","email_sent")})-SUMIF({rng("E1","mes_clave (derivada)")},"2026-08",{rng("E1","email_sent")})-SUMIF({rng("E1","mes_clave (derivada)")},"2026-09",{rng("E1","email_sent")})', S('E9b','msgs_email')-float(num('E1','email_sent')[e1mes>='2026-07'].sum()), '0.3 % por deduplicación de E1.')
# E8
res = num('E8','usuarios')
add('E8', 'E8', 'Usuarios en bloque resumen', f'=SUMIF({rng("E8","bloque")},"resumen",{rng("E8","usuarios")})', float(res[e8.bloque=='resumen'].sum()), '400,821 usuarios con baja y al menos un email antes de la baja.')
add('E8', 'E8', 'Usuarios en bloque banda_dias (debe ser igual)', f'=SUMIF({rng("E8","bloque")},"banda_dias",{rng("E8","usuarios")})', float(res[e8.bloque=='banda_dias'].sum()), '')
add('E8', 'E8', 'Usuarios en bloque banda_emails (debe ser igual)', f'=SUMIF({rng("E8","bloque")},"banda_emails",{rng("E8","usuarios")})', float(res[e8.bloque=='banda_emails'].sum()), '')
add('E8', 'E8', 'Cohorte limpia (cohorte_limpia = 1, bloque resumen)', f'=SUMIFS({rng("E8","usuarios")},{rng("E8","bloque")},"resumen",{rng("E8","cohorte_limpia")},1)', float(res[(e8.bloque=='resumen')&(e8.cohorte_limpia=='1')].sum()), 'Primer email dentro del periodo: la medida de días es exacta.')
add('E8', 'E8', 'Cohorte 0 (ya recibía antes de enero)', f'=SUMIFS({rng("E8","usuarios")},{rng("E8","bloque")},"resumen",{rng("E8","cohorte_limpia")},0)', float(res[(e8.bloque=='resumen')&(e8.cohorte_limpia=='0')].sum()), 'Días medidos desde enero: subestimados.')
add('E8', 'E8 vs E5', 'Cobertura de E8 sobre usuarios únicos con baja', f'=SUMIF({rng("E8","bloque")},"resumen",{rng("E8","usuarios")})/SUM({rng("E5","usuarios_baja")})', float(res[e8.bloque=='resumen'].sum())/S('E5','usuarios_baja'), '91 %. El resto se dio de baja sin email previo en el periodo.')
add('E8', 'E8', 'Usuarios en bloque familia_dia_baja', f'=SUMIF({rng("E8","bloque")},"familia_dia_baja",{rng("E8","usuarios")})', float(res[e8.bloque=='familia_dia_baja'].sum()), 'Mayor que 400,821 porque un usuario puede recibir varias familias el día de la baja (doble conteo ~104 %).')
# otros granos
for tab, c, ex in [('E7a','envios','E7a agrupa por familia × canal × etapa; 0.7 % debajo de E1 por filas sin familia y deduplicación. Solo para distribución por etapa.'),('E6a','msgs','Grano usuario-día. 0.5 % debajo de E1. Solo para solapamiento.'),('E7c','envios_usuario_dia','Grano usuario-día; 0.9 % debajo de E1. Solo para abren vs no abren.'),('E9a','envios_usuario_dia','Igual que E7c, con ventanas.'),('E5b','emails_despues_de_baja','1.42 M emails a usuarios ya dados de baja.'),('E6b','usuario_dias_juntas','Top 1,000 pares por LIMIT: no es el total.')]:
    add('Otros granos', tab, f'Suma de {c}', f'=SUM({rng(tab,c)})', S(tab,c), ex)
    if tab in ('E7a','E6a','E7c','E9a'):
        add('Otros granos', tab, f'  {c} menos envíos de E1', f'=SUM({rng(tab,c)})-SUM({rng("E1","envios")})', S(tab,c)-S('E1','envios'), 'Diferencia esperada por grano, ver fila anterior.')
add('E2', 'E2', 'Piezas tipo canvas', f'=COUNTIF({rng("E2","tipo_pieza")},"canvas")', int((E['E2'].tipo_pieza=='canvas').sum()), '')
add('E2', 'E2', 'Piezas tipo campaign', f'=COUNTIF({rng("E2","tipo_pieza")},"campaign")', int((E['E2'].tipo_pieza=='campaign').sum()), '')
add('E2', 'E2', 'Familias distintas en el catálogo', f'=COUNTA(UNIQUE({rng("E2","familia (derivada)")}))', int(fam_pd(E['E2'].nombre).nunique()), 'Más que en E1 porque E2 trae piezas que no enviaron en el periodo.')
add('E2', 'E2 vs E1', 'Ids de E2 que sí enviaron en el periodo (están en E1)', f'=SUMPRODUCT(--(COUNTIF({rng("E1","communication_id")},{rng("E2","communication_id")})>0))', int(E['E2'].communication_id.isin(e1.communication_id).sum()), '3,438 de 6,144. El id sin nombre de E1 no está en E2.')

# ---------- escribir Verificación
H = ['#','grupo','pestaña','qué se comprueba','valor en este Sheets (fórmula)','valor esperado (pandas sobre el CSV)','diferencia (debe ser 0)','estado','qué significa la cifra y por qué puede diferir de otra pestaña']
for j, h in enumerate(H, start=1): ver.cell(1, j, h).font = bold
for i, (g, t, q, f, exp, ex) in enumerate(V, start=2):
    ver.cell(i,1,i-1); ver.cell(i,2,g); ver.cell(i,3,t); ver.cell(i,4,q); ver.cell(i,5,f); ver.cell(i,6,exp)
    ver.cell(i,7,f'=E{i}-F{i}')
    ver.cell(i,8,f'=IF(ISERROR(E{i}),"error en fórmula",IF(ABS(G{i})<0.5,"ok","REVISAR"))')
    ver.cell(i,9,ex)
    for col in (5,6,7): ver.cell(i,col).number_format = '#,##0.00'
n = len(V)+1
ver.cell(n+2,4,'Resumen').font = bold
ver.cell(n+3,4,'Filas ok'); ver.cell(n+3,5,f'=COUNTIF(H2:H{n},"ok")')
ver.cell(n+4,4,'Filas a REVISAR'); ver.cell(n+4,5,f'=COUNTIF(H2:H{n},"REVISAR")')
ver.cell(n+5,4,'Filas con error de fórmula'); ver.cell(n+5,5,f'=COUNTIF(H2:H{n},"error en fórmula")')
ver.freeze_panes = 'A2'; ver.auto_filter.ref = f'A1:I{n}'
for col, w in zip('ABCDEFGHI', (5,14,14,62,22,22,16,20,110)): ver.column_dimensions[col].width = w

# ---------- Léeme
R = []
R.append(['Auditoría Braze Kueski: datos crudos y verificación (archivo v1, etapa 1 de 3)'])
R.append([''])
R.append(['Qué es este archivo', 'Todos los resultados de las queries de Databricks (sql/extraccion.sql, E1 a E9b) pegados completos, más los parámetros del análisis y una pestaña que comprueba que nada se perdió al pegar. En v2 se agregan Catálogo de familias y una pestaña de cálculo por lente; en v3 Estándares, Insights, Plan de acción y Vista_equipo.'])
R.append(['Fuente', 'analytics_engineering_production.xplore._dim_braze_users_and_communications (Databricks). Sin identificadores de usuario: todo viene agregado.'])
R.append(['Periodo', '1 de enero a 30 de septiembre de 2026 (event_date). E9b solo julio a septiembre.'])
R.append(['Fecha de corrida', f'Queries corridas por el usuario el {HOY}. E3, E4b, E7b y E9b en versión 2 (agregan WhatsApp entregado y quitan costo).'])
R.append(['Formato', 'Plano: encabezado en la fila 1 congelado, filtros activos, una tabla por pestaña desde A1. Columnas derivadas marcadas "(derivada)" al final de cada tabla; son fórmulas, no datos de la query.'])
R.append(['Cómo se verifica', 'Pestaña Verificación: cada fila compara un total calculado aquí con fórmula contra el mismo total calculado en pandas sobre el CSV original. La diferencia debe ser cero en todas las filas (estado ok); REVISAR significa que la pega o una fórmula falló. La última columna dice qué significa la cifra y, cuando dos pestañas no coinciden entre sí, por qué.'])
R.append([''])
R.append(['Pestaña','Archivo CSV','md5 del CSV','Filas de datos','Grano','Qué contiene','Columnas derivadas','Celdas null convertidas a vacío (columna: cuántas)'])
for tab, fname, desc, grano, _ in TABS:
    i = INFO[tab]
    nul = ', '.join(f'{c}: {n}' for c, n in i['nulls'].items() if n) or 'ninguna'
    R.append([tab, fname, i['md5'], i['filas'], grano, desc, ', '.join(i['derived']) or 'ninguna', nul])
R.append([''])
R.append(['Columnas derivadas', ''])
R.append(['familia (derivada)', 'Nombre de la comunicación sin la fecha "-AAAAMMDD-" (REGEXREPLACE). Agrupa las copias que Braze crea con fecha distinta para la misma campaña. Si la comunicación no tiene nombre, queda "(sin nombre)-" seguido del communication_id. Mismo criterio que la columna familia de E5b, E6b, E7a, E7c, E8 y E9a, que ya viene así de la query.'])
R.append(['mes_clave (derivada)', 'Primeros 7 caracteres de la fecha ("2026-06"). Las fechas llegan en formatos distintos por pestaña (E1, E4b y E6a con hora; E4a y E2 como fecha; E5 ya como mes). Todo cruce por mes usa mes_clave. En E4a las filas de grano semana llevan el mes de la fecha de inicio de la semana.'])
R.append([''])
R.append(['Definiciones que hay que tener claras', ''])
R.append(['Respuesta por canal', 'email = clic (email_click); push = apertura (pn_open); WhatsApp = lectura (whatsapp_read); in-app = clic (in_app_click). La apertura de email NO cuenta como respuesta porque Apple Mail y Gmail abren los correos automáticamente y la inflan.'])
R.append(['"Nunca respondió" (E7b y E9b, fila con resp_email, resp_push, resp_wa y resp_inapp en 0)', 'Un usuario entra aquí si en TODO el periodo (enero a septiembre en E7b; julio a septiembre en E9b) no tuvo NI UN clic en ningún email, NI UNA apertura de ningún push, NI UNA lectura de ningún WhatsApp, NI UN clic en ningún in-app, en ninguna comunicación de ningún equipo. Sí recibió mensajes: las columnas msgs_* dicen cuántos (y msgs_wa_entregados cuántos WhatsApp le llegaron de verdad). No significa que no convirtió: usuarios_con_conversion cuenta los que convirtieron aun sin responder (por ejemplo entraron solos a la app). No significa que se dio de baja: eso lo dice usuarios_con_baja_email. Una apertura de email no lo saca de este grupo.'])
R.append(['Bajas: dos cifras con nombre propio', 'CIFRA A, bajas usuario-mes = 457,376 (E4a, grano mes): un usuario cuenta en cada mes en que se dio de baja. CIFRA B, usuarios únicos con baja = 439,269 (E5 y E7b): cada usuario una vez. La diferencia (18 mil) son usuarios que se dieron de baja en más de un mes (se volvieron a suscribir o Braze registró otra baja). E1 y E3 traen EVENTOS de baja (1.09 M), que no son usuarios.'])
R.append(['Deduplicación de E1', 'E1 toma el máximo por usuario × comunicación × día, así que una fila repetida del mismo envío cuenta una vez. E3, E4b y E7b cuentan filas por usuario. Por eso E3 y E7b quedan 0.2 % arriba de E1 en email y 1.3 % en WhatsApp entregado. E1 manda en volumen por campaña; E3 y E7b mandan en distribución por usuario.'])
R.append(['usuarios_aprox', 'APPROX_COUNT_DISTINCT de Databricks: error típico de 1 a 2 %. Sumar usuarios_aprox de varias filas no da usuarios únicos.'])
R.append(['Conversión', 'Solo de referencia. conversiones_filas cuenta filas con is_converted (se repite en cada envío dentro de la ventana de atribución); conversiones_fraccional reparte cada conversión entre las comunicaciones del día; usuarios_con_conversion (E7b) cuenta usuarios. Nunca se mezclan en una misma tabla.'])
R.append(['Costo', 'Siempre con fórmula desde Parámetros: email_sent × costo por email + WhatsApp (entregado por default, o enviado) × costo por WhatsApp. Las columnas costo_mxn de E1 y E4a vienen de la query con WhatsApp sobre enviados y NO se usan. Costo fijo de Braze pendiente.'])
R.append(['Ticket', 'ticket_number viene de Braze (B0211, T2728, S0004). "Other" = sin ticket. El Catálogo (v2) lo completa con regex sobre el nombre cuando se puede.'])
R.append(['Etapa (E7a)', 'Golden record del usuario al momento del envío. Solo aparecen 3.Blacklist, 4.Kpay CC Freeze, 6.Reus or Arco, 7.Prospects y sin_golden; cubre 37 % de los envíos.'])
R.append(['E8 cohorte_limpia', '1 = el primer email del usuario cae dentro del periodo, los días hasta la baja son exactos. 0 = ya recibía antes de enero, los días se miden desde enero y están subestimados.'])
R.append([''])
R.append(['Supuestos de este archivo (aplican a todo cálculo que venga después)', ''])
for s in ['Costos: 0.05 MXN por email enviado y 1.00 MXN por WhatsApp entregado (Parámetros). Push e in-app sin costo variable. Costo fijo de Braze no incluido.',
          'WhatsApp se cobra por entregado porque Meta cobra por mensaje entregado. El selector permite ver el escenario sobre enviados.',
          'Conteos de usuarios son aproximados (APPROX_COUNT_DISTINCT); percentiles aproximados (PERCENTILE_APPROX).',
          'Familia = nombre sin fecha. Dos campañas con el mismo nombre y fecha distinta se tratan como una.',
          'Las celdas "null" de Databricks se dejaron vacías (conteo por columna en la tabla de pestañas) para que SUM y MEDIAN funcionen.',
          'E6b trae solo 1,000 pares (LIMIT en la query): sirve para ver los pares más grandes, no para totales.',
          'E8 familia_dia_baja tiene doble conteo (un usuario puede recibir varias familias el día de la baja).',
          'E1 cambia el denominador por canal: email y WhatsApp sobre entregados, push sobre enviados, in-app sobre impresiones.',
          'Equipos: E1 ya trae team mapeado (Lending aparece como Product Growth). E2 trae el team original de Braze con muchas variantes y copias.']:
    R.append(['', s])
R.append([''])
R.append(['Cómo actualizar si se vuelve a correr una query', 'Pegar el CSV completo en su pestaña desde A1 (encabezado incluido), sin borrar las columnas derivadas; arrastrar la fórmula derivada hasta la última fila; actualizar en Verificación los rangos si cambió el número de filas (los rangos están acotados a la última fila actual); anotar aquí el nuevo md5 y la fecha.'])
R.append(['Qué sigue', 'v2: Catálogo de familias y pestañas de cálculo por lente (B canales, C campañas, E después de la baja, F solapamiento, G buenas prácticas, H etapas, J canal preferido, K abren vs no abren, L días hasta la baja, M ventanas). v3: Estándares, Insights con filtros, Plan de acción y Vista_equipo.'])
for i, row in enumerate(R, start=1):
    for j, v in enumerate(row, start=1):
        c = leeme.cell(i, j, v); c.alignment = Alignment(wrap_text=True, vertical='top')
        if i == 1 or (len(row) >= 8 and j <= 8 and row[0] == 'Pestaña') or (len(row) == 2 and row[1] == ''): c.font = bold
leeme.column_dimensions['A'].width = 34; leeme.column_dimensions['B'].width = 120
for col in 'CDEFGH': leeme.column_dimensions[col].width = 30

wb.save(OUT)
print('guardado', OUT, os.path.getsize(OUT)//1024, 'KB; filas verificación', len(V))
print({t: (INFO[t]['filas'], INFO[t]['md5'][:8]) for t in INFO})
