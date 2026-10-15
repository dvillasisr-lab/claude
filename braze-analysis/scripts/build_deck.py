# -*- coding: utf-8 -*-
"""Genera el deck de la auditoría de Braze en dos formatos a partir de deck_content.py:

  deck/BRAZE_KUESKI_DECK.html   (versión web, un slide por pantalla, 16:9)
  deck/BRAZE_KUESKI_DECK.pptx   (PowerPoint con gráficas nativas)

Uso: python3 scripts/build_deck.py [--html solo] [--pptx solo]
"""
import html as H
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from deck_content import SLIDES  # noqa: E402

OUT = os.path.join(os.path.dirname(HERE), "deck")
os.makedirs(OUT, exist_ok=True)

ACCENT = "1F4FD8"
ACCENT2 = "7A9BF0"
GRAY = "A3AAB5"
GRAY2 = "D5D9E0"
INK = "1A1D24"
INK2 = "5B6270"
RED = "B5342B"
GREEN = "2C7A4B"
AMBER = "B8860B"

# ----------------------------------------------------------------------------
# utilidades
# ----------------------------------------------------------------------------

def fmtv(v, fmt):
    if v is None:
        return ""
    if fmt == "pct0":
        return f"{v:.0f} %"
    if fmt == "pct1":
        return f"{v:.1f} %"
    if fmt == "pct2":
        return f"{v:.2f} %"
    if fmt == "M":
        return f"{v:.1f} M"
    if fmt == "MXN":
        return f"{v:.2f} M"
    return f"{v}"


def status_class(s):
    s = (s or "").lower()
    if s.startswith("cumple"):
        return "ok"
    if s.startswith("debajo por poco"):
        return "warn"
    if s.startswith("debajo"):
        return "bad"
    if s in ("apagar", "confirmar apagado", "apagar y confirmar"):
        return "bad"
    if s in ("rehacer", "rehacer o re-segmentar", "bajar frecuencia"):
        return "warn"
    if s in ("proteger", "escalar"):
        return "ok"
    return ""


def section_titles():
    out = []
    for i, s in enumerate(SLIDES):
        if s["kind"] == "section":
            out.append((i, s))
    return out


# ----------------------------------------------------------------------------
# SVG
# ----------------------------------------------------------------------------

def esc(t):
    return H.escape(str(t), quote=True)


def svg_hbar(ex, W=760, highlight_color="#" + ACCENT):
    cats, vals = ex["cats"], ex["vals"]
    fmt = ex.get("fmt", "")
    notes = ex.get("notes")
    hl = set(ex.get("highlight", []))
    n = len(cats)
    rowh = 58 if n <= 5 else (48 if n <= 7 else 38)
    top = 28 if ex.get("label") else 8
    labw = 250
    noteh = 14 if notes else 0
    Hh = top + n * (rowh + noteh) + 10
    vmax = max(vals) if max(vals) > 0 else 1
    barmax = W - labw - 110
    o = [f'<svg viewBox="0 0 {W} {Hh}" class="chart" role="img" aria-label="{esc(ex.get("label",""))}">']
    if ex.get("label"):
        o.append(f'<text x="0" y="16" class="c-lab">{esc(ex["label"])}</text>')
    y = top
    for i, (c, v) in enumerate(zip(cats, vals)):
        col = highlight_color if i in hl else "#" + GRAY
        w = max(2, v / vmax * barmax)
        bh = rowh - 14
        o.append(f'<text x="{labw-10}" y="{y+bh/2+5}" text-anchor="end" class="c-cat">{esc(c)}</text>')
        o.append(f'<rect x="{labw}" y="{y}" width="{w:.1f}" height="{bh}" rx="3" fill="{col}"/>')
        o.append(f'<text x="{labw+w+8:.1f}" y="{y+bh/2+5}" class="c-val">{esc(fmtv(v, fmt))}</text>')
        if notes:
            o.append(f'<text x="{labw}" y="{y+bh+13}" class="c-note">{esc(notes[i])}</text>')
        y += rowh + noteh
    o.append("</svg>")
    return "".join(o)


def svg_hbar2(ex, W=760):
    cats, series = ex["cats"], ex["series"]
    fmt = ex.get("fmt", "")
    hl = set(ex.get("highlight", []))
    n, k = len(cats), len(series)
    bh = (18 if n <= 5 else 14) if k > 1 else 30
    gap = 3
    group = k * bh + (k - 1) * gap
    rowh = group + 18
    top = 30
    labw = 220
    Hh = top + n * rowh + 10
    vmax = ex.get("max") or max(max(s[1]) for s in series)
    barmax = W - labw - 110
    cols = ["#" + ACCENT, "#" + GRAY, "#" + GRAY2][:k] if k > 1 else None
    o = [f'<svg viewBox="0 0 {W} {Hh}" class="chart" role="img">']
    # leyenda
    x = 0
    if k > 1:
        for j, (name, _) in enumerate(series):
            o.append(f'<rect x="{x}" y="6" width="12" height="12" rx="2" fill="{cols[j]}"/><text x="{x+16}" y="17" class="c-note">{esc(name)}</text>')
            x += 16 + 7 * len(name) + 24
    else:
        o.append(f'<text x="0" y="16" class="c-lab">{esc(series[0][0])}</text>')
    y = top
    for i, c in enumerate(cats):
        o.append(f'<text x="{labw-10}" y="{y+group/2+5}" text-anchor="end" class="c-cat">{esc(c)}</text>')
        for j, (name, vals) in enumerate(series):
            v = vals[i]
            w = max(1.5, v / vmax * barmax)
            if k > 1:
                col = cols[j]
            else:
                col = "#" + ACCENT if i in hl else "#" + GRAY
            yy = y + j * (bh + gap)
            o.append(f'<rect x="{labw}" y="{yy}" width="{w:.1f}" height="{bh}" rx="2" fill="{col}"/>')
            o.append(f'<text x="{labw+w+6:.1f}" y="{yy+bh/2+4}" class="c-val{" c-sm" if k>1 else ""}">{esc(fmtv(v, fmt))}</text>')
        y += rowh
    o.append("</svg>")
    return "".join(o)


def _col_panel(o, x0, y0, w, h, cats, vals, hl, fmt, title, kind="col", ymin=None, ymax=None, ref=None):
    n = len(cats)
    o.append(f'<text x="{x0}" y="{y0-8}" class="c-lab">{esc(title)}</text>')
    lo = 0 if ymin is None else ymin
    hi = (max(vals) * 1.15) if ymax is None else ymax
    if hi <= lo:
        hi = lo + 1
    def Y(v):
        return y0 + h - (v - lo) / (hi - lo) * h
    # grid
    for t in range(0, 4):
        gv = lo + (hi - lo) * t / 3
        yy = Y(gv)
        o.append(f'<line x1="{x0}" y1="{yy:.1f}" x2="{x0+w}" y2="{yy:.1f}" class="c-grid"/>')
        o.append(f'<text x="{x0-6}" y="{yy+4:.1f}" text-anchor="end" class="c-note">{esc(fmtv(gv, fmt))}</text>')
    if ref:
        rv, rl = ref
        yy = Y(rv)
        o.append(f'<line x1="{x0}" y1="{yy:.1f}" x2="{x0+w}" y2="{yy:.1f}" class="c-ref"/><text x="{x0+w}" y="{yy-4:.1f}" text-anchor="end" class="c-note">{esc(rl)}</text>')
    slot = w / n
    if kind == "col":
        bw = slot * 0.6
        for i, (c, v) in enumerate(zip(cats, vals)):
            x = x0 + i * slot + (slot - bw) / 2
            col = "#" + ACCENT if i in hl else "#" + GRAY
            o.append(f'<rect x="{x:.1f}" y="{Y(v):.1f}" width="{bw:.1f}" height="{y0+h-Y(v):.1f}" rx="3" fill="{col}"/>')
            o.append(f'<text x="{x+bw/2:.1f}" y="{Y(v)-5:.1f}" text-anchor="middle" class="c-val c-sm">{esc(fmtv(v, fmt))}</text>')
            o.append(f'<text x="{x+bw/2:.1f}" y="{y0+h+16}" text-anchor="middle" class="c-cat">{esc(c)}</text>')
    else:
        pts = []
        for i, v in enumerate(vals):
            x = x0 + i * slot + slot / 2
            pts.append((x, Y(v)))
        d = " ".join(f'{"M" if i==0 else "L"}{x:.1f},{y:.1f}' for i, (x, y) in enumerate(pts))
        o.append(f'<path d="{d}" fill="none" stroke="#{ACCENT}" stroke-width="2.5"/>')
        for i, ((x, y), v) in enumerate(zip(pts, vals)):
            col = "#" + ACCENT
            r = 6 if i in hl else 4
            o.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="{r}" fill="{col}" stroke="#fff" stroke-width="2"/>')
            o.append(f'<text x="{x:.1f}" y="{y-10:.1f}" text-anchor="middle" class="c-val c-sm{" c-hl" if i in hl else ""}">{esc(fmtv(v, fmt))}</text>')
            o.append(f'<text x="{x:.1f}" y="{y0+h+16}" text-anchor="middle" class="c-cat">{esc(cats[i])}</text>')


def svg_colline(ex, W=760):
    cats = ex["cats"]
    hl = set(ex.get("highlight", []))
    Hh = 400
    o = [f'<svg viewBox="0 0 {W} {Hh}" class="chart" role="img">']
    name1, v1 = ex["cols"]
    name2, v2 = ex["line"]
    f1 = "M" if "(M)" in name1 else ""
    f2 = "pct2" if "%" in name2 else "pct1"
    _col_panel(o, 60, 30, W - 70, 135, cats, v1, hl, f1, name1, "col")
    _col_panel(o, 60, 236, W - 70, 135, cats, v2, hl, f2, name2, "line")
    o.append("</svg>")
    return "".join(o)


def svg_line(ex, W=760):
    cats = ex["cats"]
    hl = set(ex.get("highlight", []))
    name, vals = ex["series"][0]
    Hh = 330
    o = [f'<svg viewBox="0 0 {W} {Hh}" class="chart" role="img">']
    _col_panel(o, 60, 30, W - 70, 260, cats, vals, hl, "pct0", name, "line", ex.get("ymin"), ex.get("ymax"), ex.get("ref"))
    o.append("</svg>")
    return "".join(o)


def svg_waterfall(ex, W=760):
    cats, vals = ex["cats"], ex["vals"]
    fmt = ex.get("fmt", "")
    n = len(cats)
    Hh = 360
    x0, y0, w, h = 60, 40, W - 80, 240
    o = [f'<svg viewBox="0 0 {W} {Hh}" class="chart" role="img">']
    o.append(f'<text x="{x0}" y="22" class="c-lab">{esc(ex.get("label",""))}</text>')
    hi = max(vals) * 1.1
    def Y(v):
        return y0 + h - v / hi * h
    slot = w / n
    bw = slot * 0.55
    run = 0
    for i, (c, v) in enumerate(zip(cats, vals)):
        x = x0 + i * slot + (slot - bw) / 2
        if i == 0 or i == n - 1:
            top, bot, col = v, 0, "#" + (ACCENT if i == n - 1 else INK2)
            run = v
        else:
            top, bot = run, run + v
            col = "#" + RED
            run = run + v
            o.append(f'<line x1="{x-slot*0.225:.1f}" y1="{Y(top):.1f}" x2="{x:.1f}" y2="{Y(top):.1f}" class="c-grid"/>')
        yt, yb = Y(max(top, bot)), Y(min(top, bot))
        o.append(f'<rect x="{x:.1f}" y="{yt:.1f}" width="{bw:.1f}" height="{max(2, yb-yt):.1f}" rx="3" fill="{col}"/>')
        lab = fmtv(abs(v), fmt) if i not in (0, n - 1) else fmtv(v, fmt)
        if i not in (0, n - 1):
            lab = "−" + lab
        o.append(f'<text x="{x+bw/2:.1f}" y="{yt-6:.1f}" text-anchor="middle" class="c-val c-sm">{esc(lab)}</text>')
        # categoría en dos líneas si es larga
        parts = c.split(" (")
        o.append(f'<text x="{x+bw/2:.1f}" y="{y0+h+18}" text-anchor="middle" class="c-cat">{esc(parts[0])}</text>')
        if len(parts) > 1:
            o.append(f'<text x="{x+bw/2:.1f}" y="{y0+h+34}" text-anchor="middle" class="c-note">({esc(parts[1])}</text>')
    o.append(f'<text x="{x0+w}" y="{y0+h+60}" text-anchor="end" class="c-note">Ahorro total: 1.60 M MXN al mes (33 %). Confirmar apagado (232) y bajar frecuencia (84) no suman costo hoy.</text>')
    o.append("</svg>")
    return "".join(o)


def html_table(ex):
    head, rows = ex["head"], ex["rows"]
    sc = ex.get("status_col")
    scs = sc if isinstance(sc, list) else ([sc] if sc is not None else [])
    hr = set(ex.get("highlight_rows", []))
    cls = "tbl" + (" tbl-sm" if (ex.get("small") or len(rows) >= 8) else "")
    o = [f'<div class="tblwrap"><table class="{cls}"><thead><tr>']
    for h in head:
        o.append(f"<th>{esc(h)}</th>")
    o.append("</tr></thead><tbody>")
    for ri, r in enumerate(rows):
        o.append(f'<tr{" class=\"hl\"" if ri in hr else ""}>')
        for ci, c in enumerate(r):
            if ci in scs:
                k = status_class(c)
                o.append(f'<td><span class="pill {k}">{esc(c)}</span></td>')
            else:
                o.append(f"<td>{esc(c)}</td>")
        o.append("</tr>")
    o.append("</tbody></table></div>")
    return "".join(o)


def html_stats(ex):
    o = ['<div class="stats">']
    for big, lab, sub in ex["items"]:
        o.append(f'<div class="stat"><div class="s-big">{esc(big)}</div><div class="s-lab">{esc(lab)}</div><div class="s-sub">{esc(sub)}</div></div>')
    o.append("</div>")
    return "".join(o)


def html_rules(ex):
    o = ['<ol class="rules">']
    for r in ex["items"]:
        o.append(f"<li>{esc(r)}</li>")
    o.append("</ol>")
    return "".join(o)


def html_gantt(ex):
    cols = ex["cols"]
    o = ['<div class="tblwrap"><table class="tbl tbl-sm gantt"><thead><tr><th>Qué</th>']
    for c in cols:
        o.append(f"<th class=\"gc\">{esc(c)}</th>")
    o.append("<th>Quién</th></tr></thead><tbody>")
    for name, cells, who in ex["rows"]:
        o.append(f"<tr><td>{esc(name)}</td>")
        for v in cells:
            o.append(f'<td class="gc"><span class="gbar{" on" if v else ""}"></span></td>')
        o.append(f"<td class=\"who\">{esc(who)}</td></tr>")
    o.append("</tbody></table></div>")
    return "".join(o)


def render_exhibit(ex):
    t = ex["type"]
    if t == "hbar":
        return svg_hbar(ex)
    if t == "hbar2":
        return svg_hbar2(ex)
    if t == "colline":
        return svg_colline(ex)
    if t == "line":
        return svg_line(ex)
    if t == "waterfall":
        return svg_waterfall(ex)
    if t == "table":
        return html_table(ex)
    if t == "stats":
        return html_stats(ex)
    if t == "rules":
        return html_rules(ex)
    if t == "gantt":
        return html_gantt(ex)
    raise ValueError(t)


# ----------------------------------------------------------------------------
# HTML
# ----------------------------------------------------------------------------

CSS = """
/* Layout: un slide 16:9 por pantalla, escalado con container queries; chrome exterior sigue el tema. */
:root{
  --accent:#%(A)s; --accent-soft:#E4EAFB; --ink:#%(I)s; --ink-2:#%(I2)s; --ink-3:#8A909B;
  --rule:#E3E6EB; --paper:#FFFFFF; --paper-2:#F6F7F9; --chrome:#EEF0F3; --chrome-ink:#4B515C;
  --bad:#%(R)s; --ok:#%(G)s; --warn:#%(AM)s; --bad-bg:#F8E7E5; --ok-bg:#E3F1E8; --warn-bg:#F7EFD6;
  --font:"Source Sans 3","Segoe UI",system-ui,sans-serif;
}
@media (prefers-color-scheme: dark){ :root:not([data-theme="light"]){ --chrome:#15171B; --chrome-ink:#B9BEC7; --accent-soft:#E4EAFB; color-scheme:dark; } }
:root[data-theme="dark"]{ --chrome:#15171B; --chrome-ink:#B9BEC7; --accent-soft:#E4EAFB; color-scheme:dark; }
*{box-sizing:border-box}
body{margin:0;background:var(--chrome);color:var(--ink);font-family:var(--font);padding-inline:16px;padding-block:12px 40px}
.deck{max-width:1280px;margin:0 auto;display:flex;flex-direction:column;gap:18px}
.topbar{display:flex;flex-wrap:wrap;gap:8px 16px;align-items:center;color:var(--chrome-ink);font-size:14px;max-width:1280px;margin:0 auto 10px}
.topbar strong{color:var(--chrome-ink)}
.topbar .sp{flex:1}
.topbar a{color:var(--chrome-ink)}
.toc{display:flex;flex-wrap:wrap;gap:6px}
.toc a{display:inline-block;padding:3px 10px;border:1px solid currentColor;border-radius:999px;text-decoration:none;font-size:12.5px;opacity:.85}
.toc a:hover,.toc a:focus-visible{opacity:1;outline:2px solid var(--accent);outline-offset:1px}
.slide{container-type:inline-size;aspect-ratio:16/9;width:100%%;background:var(--paper);color:var(--ink);border-radius:6px;box-shadow:0 1px 2px rgba(0,0,0,.08),0 6px 24px rgba(0,0,0,.06);position:relative;overflow:hidden;scroll-margin-top:12px}
.s{position:absolute;inset:0;padding:4.2cqi 4.6cqi 3.4cqi;display:flex;flex-direction:column;gap:1.6cqi}
.tracker{position:absolute;top:2.2cqi;right:4.6cqi;font-size:1.25cqi;letter-spacing:.06em;text-transform:uppercase;color:var(--ink-3)}
.num{position:absolute;bottom:1.8cqi;right:4.6cqi;font-size:1.25cqi;color:var(--ink-3);font-variant-numeric:tabular-nums}
h2.t{margin:0;font-size:2.75cqi;line-height:1.2;font-weight:600;letter-spacing:-.01em;text-wrap:balance;max-width:84cqi;color:var(--ink)}
h2.t::after{content:"";display:block;width:6cqi;height:.35cqi;background:var(--accent);margin-top:1.2cqi;border-radius:2px}
.body{flex:1;display:grid;grid-template-columns:minmax(0,1.55fr) minmax(0,1fr);gap:3.4cqi;min-height:0}
.body.full{grid-template-columns:minmax(0,1fr)}
.ex{min-width:0;min-height:0;font-size:1.5cqi}
.side{min-width:0;font-size:1.55cqi;line-height:1.4;color:var(--ink-2);display:flex;flex-direction:column;gap:1cqi}
.side p{margin:0;padding-left:1.4cqi;border-left:.3cqi solid var(--accent-soft)}
.side p strong{color:var(--ink)}
.src{font-size:1.15cqi;color:var(--ink-3);border-top:1px solid var(--rule);padding-top:.8cqi;padding-right:9cqi}
.chart{width:100%%;height:auto;display:block;max-height:min(100%%,40cqi)}
.sowhat{display:flex;gap:1.4cqi;align-items:flex-start;background:var(--accent-soft);border-left:.4cqi solid var(--accent);padding:1cqi 1.4cqi;margin-top:1.2cqi;font-size:1.5cqi;line-height:1.35;color:var(--ink)}
.sowhat b{flex:0 0 auto;color:var(--accent);font-size:1.1cqi;letter-spacing:.1em;text-transform:uppercase;padding-top:.3cqi}
.c-lab{font-size:15px;font-weight:600;fill:var(--ink)}
.c-cat{font-size:13.5px;fill:var(--ink-2)}
.c-val{font-size:13.5px;font-weight:600;fill:var(--ink);font-variant-numeric:tabular-nums}
.c-val.c-sm{font-size:12px}
.c-val.c-hl{fill:var(--accent)}
.c-note{font-size:11px;fill:var(--ink-3)}
.c-grid{stroke:var(--rule);stroke-width:1}
.c-ref{stroke:var(--ink-3);stroke-width:1;stroke-dasharray:4 4}
.tblwrap{overflow:auto;max-height:50cqi}
.tbl{width:100%%;border-collapse:collapse;font-size:1.35cqi;line-height:1.3}
.tbl th{text-align:left;font-weight:600;color:var(--ink-2);font-size:1.15cqi;letter-spacing:.04em;text-transform:uppercase;border-bottom:.18cqi solid var(--ink);padding:.5cqi .8cqi .5cqi 0}
.tbl td{padding:.62cqi .8cqi .62cqi 0;border-bottom:1px solid var(--rule);vertical-align:top;color:var(--ink)}
.tbl td:first-child{font-weight:600}
.tbl tr.hl td{background:var(--accent-soft)}
.tbl-sm{font-size:1.15cqi}
.tbl-sm td,.tbl-sm th{padding:.38cqi .5cqi .38cqi 0}
.pill{display:inline-block;padding:.1cqi .7cqi;border-radius:999px;font-size:1.1cqi;font-weight:600;white-space:nowrap;background:var(--paper-2);color:var(--ink-2)}
.pill.ok{background:var(--ok-bg);color:var(--ok)}.pill.bad{background:var(--bad-bg);color:var(--bad)}.pill.warn{background:var(--warn-bg);color:var(--warn)}
.stats{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:1.6cqi 2.4cqi;align-content:start}
.stats:has(.stat:nth-child(5)){grid-template-columns:repeat(3,minmax(0,1fr))}
.stat{border-top:.25cqi solid var(--accent);padding-top:1cqi}
.s-big{font-size:3.6cqi;font-weight:700;letter-spacing:-.02em;line-height:1;color:var(--accent);font-variant-numeric:tabular-nums}
.s-lab{font-size:1.5cqi;font-weight:600;margin-top:.6cqi;color:var(--ink)}
.s-sub{font-size:1.25cqi;color:var(--ink-2);margin-top:.3cqi}
.rules{margin:0;padding:0;list-style:none;counter-reset:r;columns:2;column-gap:3cqi;font-size:1.42cqi;line-height:1.35}
.rules li{counter-increment:r;break-inside:avoid;display:flex;gap:.9cqi;margin-bottom:1cqi}
.rules li::before{content:counter(r);flex:0 0 2.4cqi;height:2.4cqi;border-radius:50%%;background:var(--accent);color:#fff;font-weight:700;font-size:1.25cqi;display:inline-flex;align-items:center;justify-content:center}
.gantt .gc{width:10cqi;text-align:center}
.gbar{display:block;height:1.4cqi;border-radius:3px;background:transparent}
.gbar.on{background:var(--accent)}
.gantt .who{color:var(--ink-2);font-weight:400}
/* portada y divisores */
.slide.cover .s{justify-content:center;background:linear-gradient(135deg,#FFFFFF 0%%,#EEF2FC 100%%)}
.cover h1{font-size:5.2cqi;line-height:1.08;margin:0;font-weight:700;letter-spacing:-.02em;max-width:70cqi;text-wrap:balance}
.cover .sub{font-size:2.2cqi;color:var(--ink-2);margin-top:1.6cqi}
.cover .meta{font-size:1.5cqi;color:var(--ink-3);margin-top:4cqi}
.cover .brand{position:absolute;top:4.2cqi;left:4.6cqi;font-size:1.3cqi;letter-spacing:.12em;text-transform:uppercase;color:var(--accent);font-weight:700}
.slide.sec .s{justify-content:center;background:var(--accent);color:#fff}
.sec .big{font-size:9cqi;font-weight:800;line-height:1;opacity:.35}
.sec h2{font-size:4.2cqi;margin:.6cqi 0 0;font-weight:700;letter-spacing:-.015em;max-width:70cqi;text-wrap:balance}
.sec p{font-size:1.9cqi;opacity:.9;margin:1.2cqi 0 0;max-width:70cqi}
/* resumen ejecutivo */
.exec{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:3cqi;flex:1}
.exec h3{margin:0 0 1cqi;font-size:1.3cqi;letter-spacing:.08em;text-transform:uppercase;color:var(--accent)}
.exec ul{margin:0;padding:0;list-style:none;display:flex;flex-direction:column;gap:1.2cqi;font-size:1.45cqi;line-height:1.4;color:var(--ink)}
.exec li{padding-left:1.2cqi;border-left:.3cqi solid var(--accent-soft)}
@media (max-width:700px){ .body{grid-template-columns:minmax(0,1fr)} .side{display:none} .exec{grid-template-columns:minmax(0,1fr)} .rules{columns:1} }
@media print{ body{background:#fff;padding:0} .topbar{display:none} .slide{box-shadow:none;border-radius:0;page-break-after:always;width:100%%} .side{display:flex} }
""" % {"A": ACCENT, "I": INK, "I2": INK2, "R": RED, "G": GREEN, "AM": AMBER}

JS = """
(function(){
  var slides=[].slice.call(document.querySelectorAll('.slide'));
  var cur=0;
  function idx(){var y=window.scrollY+window.innerHeight/2;for(var i=slides.length-1;i>=0;i--){if(slides[i].offsetTop<=y)return i;}return 0;}
  function go(i){i=Math.max(0,Math.min(slides.length-1,i));slides[i].scrollIntoView({behavior:'smooth',block:'start'});}
  document.addEventListener('keydown',function(e){
    if(e.target.tagName==='INPUT')return;
    if(e.key==='ArrowRight'||e.key==='PageDown'||e.key===' '){e.preventDefault();go(idx()+1);}
    if(e.key==='ArrowLeft'||e.key==='PageUp'){e.preventDefault();go(idx()-1);}
    if(e.key==='Home'){go(0);} if(e.key==='End'){go(slides.length-1);}
  });
  var pos=document.getElementById('pos');
  function upd(){var i=idx();if(i!==cur){cur=i;}pos.textContent=(i+1)+' / '+slides.length;}
  window.addEventListener('scroll',upd,{passive:true});upd();
})();
"""


def sowhat(s):
    t = s.get("sowhat")
    return f'<div class="sowhat"><b>So what</b><span>{esc(t)}</span></div>' if t else ""


def build_html():
    o = []
    o.append("<title>Braze en Kueski</title>")
    o.append('<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Source+Sans+3:wght@400;600;700;800&display=swap">')
    o.append(f"<style>{CSS}</style>")
    n = len(SLIDES)
    toc = "".join(f'<a href="#s{i+1}">{esc(s["num"])}. {esc(s["title"])}</a>' for i, s in section_titles())
    o.append(f'<div class="topbar"><strong>Braze en Kueski</strong><span>Auditoría enero a septiembre de 2026</span><span class="sp"></span><span class="toc"><a href="#s1">Inicio</a><a href="#s2">Resumen</a>{toc}</span><span id="pos">1 / {n}</span></div>')
    o.append('<div class="deck">')
    for i, s in enumerate(SLIDES, 1):
        k = s["kind"]
        if k == "title":
            o.append(f'<section class="slide cover" id="s{i}"><div class="s"><div class="brand">Kueski · Braze</div><h1>{esc(s["title"])}</h1><div class="sub">{esc(s["subtitle"])}</div><div class="meta">{esc(s["meta"])}</div></div><div class="num">{i}</div></section>')
        elif k == "section":
            o.append(f'<section class="slide sec" id="s{i}"><div class="s"><div class="big">{esc(s["num"])}</div><h2>{esc(s["title"])}</h2>{"<p>"+esc(s["subtitle"])+"</p>" if s.get("subtitle") else ""}</div><div class="num" style="color:#fff;opacity:.7">{i}</div></section>')
        elif k == "exec":
            cols = "".join(f'<div><h3>{esc(h)}</h3><ul>{"".join("<li>"+esc(b)+"</li>" for b in bs)}</ul></div>' for h, bs in s["groups"])
            o.append(f'<section class="slide" id="s{i}"><div class="s"><div class="tracker">Resumen ejecutivo</div><h2 class="t">{esc(s["title"])}</h2><div class="exec">{cols}</div>{sowhat(s)}<div class="src">{esc("Fuente: Sheets braze_auditoria_v4 (Insights, Plan_accion, Estándares); análisis")}</div></div><div class="num">{i}</div></section>')
        else:
            side = s.get("side") or []
            ex = render_exhibit(s["exhibit"])
            sidehtml = "".join(f"<p>{esc(b)}</p>" for b in side)
            body = f'<div class="body{"" if side else " full"}"><div class="ex">{ex}</div>{"<div class=\"side\">"+sidehtml+"</div>" if side else ""}</div>'
            o.append(f'<section class="slide" id="s{i}"><div class="s"><div class="tracker">{esc(s.get("tracker",""))}</div><h2 class="t">{esc(s["title"])}</h2>{body}{sowhat(s)}<div class="src">{esc(s.get("source",""))}</div></div><div class="num">{i}</div></section>')
    o.append("</div>")
    o.append(f"<script>{JS}</script>")
    path = os.path.join(OUT, "BRAZE_KUESKI_DECK.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(o))
    return path


# ----------------------------------------------------------------------------
# PPTX
# ----------------------------------------------------------------------------

def build_pptx():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt, Emu

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    C = lambda h: RGBColor.from_string(h)
    FONT = "Calibri"

    def tb(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT):
        s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = s.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = anchor
        lines = text if isinstance(text, list) else [text]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ln
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = font
            r.font.color.rgb = C(color)
        return s

    def bullets(slide, x, y, w, h, items, size=12, color=INK2):
        s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = s.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            r = p.add_run()
            r.text = "•  " + it
            r.font.size = Pt(size)
            r.font.name = FONT
            r.font.color.rgb = C(color)
        return s

    def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = C(color)
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def sowhat_box(slide, s):
        t = s.get("sowhat")
        if not t:
            return
        rect(slide, 0.6, 6.1, 12.13, 0.72, "E4EAFB")
        rect(slide, 0.6, 6.1, 0.06, 0.72, ACCENT)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.8), Inches(0.72))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.05)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = "SO WHAT   "
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = C(ACCENT)
        r = p.add_run()
        r.text = t
        r.font.size = Pt(11)
        r.font.name = FONT
        r.font.color.rgb = C(INK)

    def chrome(slide, num, tracker, title, source):
        if tracker:
            tb(slide, 9.0, 0.25, 3.9, 0.3, tracker.upper(), 9, False, "8A909B", PP_ALIGN.RIGHT)
        tb(slide, 0.6, 0.5, 11.5, 1.1, title, 22, True, INK, anchor=MSO_ANCHOR.TOP)
        rect(slide, 0.6, 1.62, 0.8, 0.05, ACCENT)
        if source:
            ln = slide.shapes.add_connector(1, Inches(0.6), Inches(6.95), Inches(12.73), Inches(6.95))
            ln.line.color.rgb = C("E3E6EB")
            tb(slide, 0.6, 7.0, 11.2, 0.35, source, 8.5, False, "8A909B")
        tb(slide, 12.3, 7.0, 0.45, 0.35, str(num), 9, False, "8A909B", PP_ALIGN.RIGHT)

    def numfmt(fmt):
        return {"pct0": '0" %"', "pct1": '0.0" %"', "pct2": '0.00" %"', "M": '0.0" M"', "MXN": '0.00" M"'}.get(fmt, "General")

    def style_axis(ax, size=10, visible=True):
        ax.tick_labels.font.size = Pt(size)
        ax.tick_labels.font.name = FONT
        ax.tick_labels.font.color.rgb = C(INK2)
        ax.format.line.color.rgb = C("E3E6EB")
        ax.has_major_gridlines = False
        if not visible:
            ax.visible = False

    def add_chart(slide, kind, x, y, w, h, cats, series, fmt, highlight=(), legend=False, title=None, ymin=None, ymax=None):
        cd = CategoryChartData()
        cd.categories = cats
        for name, vals in series:
            cd.add_series(name, vals)
        gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.font.name = FONT
        ch.font.size = Pt(10)
        ch.has_legend = legend
        if legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(9)
        if title:
            ch.has_title = True
            ch.chart_title.text_frame.text = title
            p = ch.chart_title.text_frame.paragraphs[0]
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.bold = True
            p.runs[0].font.name = FONT
            p.runs[0].font.color.rgb = C(INK)
        else:
            ch.has_title = False
        plot = ch.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = numfmt(fmt)
        dl.number_format_is_linked = False
        dl.font.size = Pt(9)
        dl.font.name = FONT
        dl.font.color.rgb = C(INK)
        try:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END if kind != XL_CHART_TYPE.LINE_MARKERS else XL_LABEL_POSITION.ABOVE
        except Exception:
            pass
        if kind in (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.COLUMN_CLUSTERED):
            plot.gap_width = 60 if len(series) == 1 else 80
            plot.overlap = -10 if len(series) > 1 else 0
        style_axis(ch.category_axis, 10)
        va = ch.value_axis
        style_axis(va, 9, visible=False)
        va.has_major_gridlines = False
        if ymin is not None:
            va.minimum_scale = ymin
        if ymax is not None:
            va.maximum_scale = ymax
        if kind == XL_CHART_TYPE.BAR_CLUSTERED:
            ch.category_axis.reverse_order = True
            ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        pal = [ACCENT, GRAY, GRAY2]
        for si, s in enumerate(plot.series):
            if kind == XL_CHART_TYPE.LINE_MARKERS:
                s.format.line.color.rgb = C(ACCENT)
                s.format.line.width = Pt(2.25)
                s.marker.style = 8
                s.marker.size = 7
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = C(ACCENT)
                s.marker.format.line.color.rgb = C("FFFFFF")
                s.smooth = False
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = C(pal[si] if len(series) > 1 else GRAY)
                s.invert_if_negative = False
                if len(series) == 1:
                    for pi in range(len(cats)):
                        pt = s.points[pi]
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = C(ACCENT if pi in highlight else GRAY)
        return ch

    def add_table(slide, ex, x, y, w, h):
        head, rows = ex["head"], ex["rows"]
        sc = ex.get("status_col")
        scs = sc if isinstance(sc, list) else ([sc] if sc is not None else [])
        hr = set(ex.get("highlight_rows", []))
        nr, nc = len(rows) + 1, len(head)
        shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h))
        t = shp.table
        small = ex.get("small") or nr > 11 or nc > 6
        fs = 8 if small else (9 if nr >= 8 else 10)
        # anchos: primera columna más ancha cuando hay texto largo
        lens = [max(len(str(r[c])) for r in [head] + rows) for c in range(nc)]
        tot = sum(min(l, 60) for l in lens)
        for c in range(nc):
            t.columns[c].width = Emu(int(Inches(w) * min(lens[c], 60) / tot))
        for ri in range(nr):
            t.rows[ri].height = Inches(0.28)
            for ci in range(nc):
                cell = t.cell(ri, ci)
                txt = head[ci] if ri == 0 else str(rows[ri - 1][ci])
                cell.margin_left = cell.margin_right = Inches(0.05)
                cell.margin_top = cell.margin_bottom = Inches(0.02)
                cell.fill.solid()
                if ri == 0:
                    cell.fill.fore_color.rgb = C("FFFFFF")
                elif (ri - 1) in hr:
                    cell.fill.fore_color.rgb = C("E4EAFB")
                else:
                    cell.fill.fore_color.rgb = C("FFFFFF" if ri % 2 else "F6F7F9")
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = txt
                r.font.name = FONT
                r.font.size = Pt(fs if ri else fs - 1)
                r.font.bold = (ri == 0) or (ci == 0)
                col = INK2 if ri == 0 else INK
                if ri and ci in scs:
                    k = status_class(txt)
                    col = {"ok": GREEN, "bad": RED, "warn": AMBER}.get(k, INK)
                    r.font.bold = True
                r.font.color.rgb = C(col)
        return shp

    def add_stats(slide, ex, x, y, w, h):
        items = ex["items"]
        cols = 3 if len(items) > 4 else 2
        rows = (len(items) + cols - 1) // cols
        cw, rh = w / cols, h / rows
        for i, (big, lab, sub) in enumerate(items):
            cx, cy = x + (i % cols) * cw, y + (i // cols) * rh
            rect(slide, cx, cy, cw - 0.3, 0.04, ACCENT)
            tb(slide, cx, cy + 0.12, cw - 0.3, 0.6, big, 26, True, ACCENT)
            tb(slide, cx, cy + 0.78, cw - 0.3, 0.4, lab, 12, True, INK)
            tb(slide, cx, cy + 1.12, cw - 0.3, 0.6, sub, 9.5, False, INK2)

    def add_rules(slide, ex, x, y, w, h):
        items = ex["items"]
        half = (len(items) + 1) // 2
        for col in range(2):
            sub = items[col * half:(col + 1) * half]
            s = slide.shapes.add_textbox(Inches(x + col * w / 2), Inches(y), Inches(w / 2 - 0.3), Inches(h))
            tf = s.text_frame
            tf.word_wrap = True
            for i, it in enumerate(sub):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(7)
                r = p.add_run()
                r.text = f"{col*half+i+1}.  "
                r.font.bold = True
                r.font.size = Pt(11)
                r.font.name = FONT
                r.font.color.rgb = C(ACCENT)
                r2 = p.add_run()
                r2.text = it
                r2.font.size = Pt(11)
                r2.font.name = FONT
                r2.font.color.rgb = C(INK)

    def add_gantt(slide, ex, x, y, w, h):
        cols = ex["cols"]
        rows = ex["rows"]
        name_w, who_w = 4.6, 2.6
        cw = (w - name_w - who_w) / len(cols)
        rh = min(0.42, (h - 0.4) / len(rows))
        for j, c in enumerate(cols):
            tb(slide, x + name_w + j * cw, y, cw, 0.3, c, 9, True, INK2, PP_ALIGN.CENTER)
        ln = slide.shapes.add_connector(1, Inches(x), Inches(y + 0.32), Inches(x + w), Inches(y + 0.32))
        ln.line.color.rgb = C(INK)
        for i, (name, cells, who) in enumerate(rows):
            yy = y + 0.4 + i * rh
            tb(slide, x, yy, name_w - 0.1, rh, name, 9.5, True, INK, anchor=MSO_ANCHOR.MIDDLE)
            for j, v in enumerate(cells):
                if v:
                    rect(slide, x + name_w + j * cw + 0.08, yy + rh * 0.3, cw - 0.16, rh * 0.4, ACCENT)
            tb(slide, x + name_w + len(cols) * cw + 0.1, yy, who_w - 0.1, rh, who, 9, False, INK2, anchor=MSO_ANCHOR.MIDDLE)
            ln = slide.shapes.add_connector(1, Inches(x), Inches(yy + rh), Inches(x + w), Inches(yy + rh))
            ln.line.color.rgb = C("E3E6EB")

    def add_waterfall(slide, ex, x, y, w, h):
        cats, vals = ex["cats"], ex["vals"]
        n = len(cats)
        base, delta, ends = [], [], []
        run = 0
        for i, v in enumerate(vals):
            if i == 0 or i == n - 1:
                base.append(0); delta.append(0); ends.append(v); run = v
            else:
                run2 = run + v
                base.append(min(run, run2)); delta.append(abs(v)); ends.append(0); run = run2
        cd = CategoryChartData()
        cd.categories = cats
        cd.add_series("base", base)
        cd.add_series("Totales", ends)
        cd.add_series("Ahorro", delta)
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.font.name = FONT
        ch.font.size = Pt(10)
        ch.has_legend = False
        ch.has_title = True
        ch.chart_title.text_frame.text = ex.get("label", "")
        r0 = ch.chart_title.text_frame.paragraphs[0].runs[0]
        r0.font.size = Pt(11); r0.font.bold = True; r0.font.name = FONT; r0.font.color.rgb = C(INK)
        plot = ch.plots[0]
        plot.gap_width = 55
        style_axis(ch.category_axis, 9)
        style_axis(ch.value_axis, 9, visible=False)
        s0, s1, s2 = plot.series
        s0.format.fill.background()
        s1.format.fill.solid(); s1.format.fill.fore_color.rgb = C(INK2)
        s1.points[n - 1].format.fill.solid(); s1.points[n - 1].format.fill.fore_color.rgb = C(ACCENT)
        s2.format.fill.solid(); s2.format.fill.fore_color.rgb = C(RED)
        for s, nf in ((s1, '0.00" M";;'), (s2, '"−"0.00" M";;')):
            s.data_labels.show_value = True
            s.data_labels.number_format = nf
            s.data_labels.number_format_is_linked = False
            s.data_labels.font.size = Pt(9)
            s.data_labels.font.name = FONT
            s.data_labels.font.color.rgb = C(INK)
            try:
                s.data_labels.position = XL_LABEL_POSITION.INSIDE_END
            except Exception:
                pass
        tb(slide, x, y + h, w, 0.3, "Ahorro total: 1.60 M MXN al mes (33 %). Confirmar apagado (232) y bajar frecuencia (84) no suman costo hoy.", 9, False, INK2)

    def exhibit(slide, ex, x, y, w, h):
        t = ex["type"]
        fmt = ex.get("fmt", "")
        hl = set(ex.get("highlight", []))
        if t == "hbar":
            notes = ex.get("notes")
            cats = ex["cats"] if not notes else [f"{c}  ({n})" if n else c for c, n in zip(ex["cats"], notes)]
            add_chart(slide, XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, cats, [(ex.get("label", ""), ex["vals"])], fmt, hl, title=ex.get("label"))
        elif t == "hbar2":
            add_chart(slide, XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, ex["cats"], ex["series"], fmt, hl, legend=len(ex["series"]) > 1, title=None if len(ex["series"]) > 1 else ex["series"][0][0], ymax=ex.get("max"))
        elif t == "colline":
            n1, v1 = ex["cols"]; n2, v2 = ex["line"]
            add_chart(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h / 2 - 0.05, ex["cats"], [(n1, v1)], "M" if "(M)" in n1 else "", hl, title=n1)
            add_chart(slide, XL_CHART_TYPE.LINE_MARKERS, x, y + h / 2 + 0.05, w, h / 2 - 0.05, ex["cats"], [(n2, v2)], "pct2" if "%" in n2 else "pct1", hl, title=n2)
        elif t == "line":
            n1, v1 = ex["series"][0]
            add_chart(slide, XL_CHART_TYPE.LINE_MARKERS, x, y, w, h, ex["cats"], [(n1, v1)], "pct0", hl, title=n1, ymin=ex.get("ymin"), ymax=ex.get("ymax"))
            if ex.get("ref"):
                tb(slide, x, y + h, w, 0.3, f"Línea de referencia: {ex['ref'][1]}", 9, False, INK2)
        elif t == "waterfall":
            add_waterfall(slide, ex, x, y, w, h - 0.3)
        elif t == "table":
            add_table(slide, ex, x, y, w, h)
        elif t == "stats":
            add_stats(slide, ex, x, y, w, h)
        elif t == "rules":
            add_rules(slide, ex, x, y, w, h)
        elif t == "gantt":
            add_gantt(slide, ex, x, y, w, h)

    for i, s in enumerate(SLIDES, 1):
        sl = prs.slides.add_slide(blank)
        k = s["kind"]
        if k == "title":
            rect(sl, 0, 0, 13.333, 7.5, "FFFFFF")
            rect(sl, 0, 0, 0.25, 7.5, ACCENT)
            tb(sl, 0.9, 0.6, 6, 0.4, "KUESKI · BRAZE", 11, True, ACCENT)
            tb(sl, 0.9, 2.3, 9.5, 1.8, s["title"], 40, True, INK)
            tb(sl, 0.9, 4.2, 9.5, 0.6, s["subtitle"], 18, False, INK2)
            tb(sl, 0.9, 5.6, 9.5, 0.4, s["meta"], 12, False, "8A909B")
        elif k == "section":
            rect(sl, 0, 0, 13.333, 7.5, ACCENT)
            tb(sl, 0.9, 1.6, 6, 1.4, s["num"], 72, True, "FFFFFF")
            tb(sl, 0.9, 3.2, 10, 1.2, s["title"], 32, True, "FFFFFF")
            if s.get("subtitle"):
                tb(sl, 0.9, 4.4, 10, 0.8, s["subtitle"], 16, False, "FFFFFF")
            tb(sl, 12.3, 7.0, 0.45, 0.35, str(i), 9, False, "FFFFFF", PP_ALIGN.RIGHT)
        elif k == "exec":
            chrome(sl, i, "Resumen ejecutivo", s["title"], "Fuente: Sheets braze_auditoria_v4 (Insights, Plan_accion, Estándares); análisis")
            cw = 12.1 / 3
            for j, (h, bs) in enumerate(s["groups"]):
                x = 0.6 + j * cw
                tb(sl, x, 1.95, cw - 0.3, 0.3, h.upper(), 10, True, ACCENT)
                bullets(sl, x, 2.3, cw - 0.3, 3.7, bs, 11.5, INK)
            sowhat_box(sl, s)
        else:
            chrome(sl, i, s.get("tracker", ""), s["title"], s.get("source", ""))
            side = s.get("side") or []
            if side:
                exhibit(sl, s["exhibit"], 0.6, 1.95, 7.4, 4.05)
                bullets(sl, 8.3, 1.95, 4.45, 4.05, side, 10.5, INK2)
            else:
                exhibit(sl, s["exhibit"], 0.6, 1.95, 12.13, 4.05)
            sowhat_box(sl, s)
    path = os.path.join(OUT, "BRAZE_KUESKI_DECK.pptx")
    prs.save(path)
    return path


if __name__ == "__main__":
    args = sys.argv[1:]
    if "--pptx" not in args:
        print("html:", build_html())
    if "--html" not in args:
        print("pptx:", build_pptx())
